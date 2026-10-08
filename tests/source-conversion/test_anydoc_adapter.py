"""Behavior contracts plus actual pinned native backend where locally installed."""
import argparse
import hashlib
import importlib.util
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location("anydoc_source_converter", ROOT / "tools/source-to-markdown.py")
C = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = C
SPEC.loader.exec_module(C)
SITE = ROOT / "wiki/_outputs/source-conversions/runtimes/anydoc-0.2.4/site"


class AnyDocTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="anydoc-contract-")
        self.root = Path(self.temp.name)
        (self.root / "raw/assets").mkdir(parents=True)

    def tearDown(self):
        self.temp.cleanup()

    def args(self, **changes):
        args = vars(C.parse_args([]))
        args.update(backend="anydoc", anydoc_site=str(SITE))
        args.update(changes)
        return argparse.Namespace(**args)

    def item(self, path):
        return C.InventoryItem(path, "raw", Path(path.name), Path("raw/assets") / path.name,
            path.suffix, path.stat().st_size, "2026-09-11", "document", "convert",
            f"wiki/_extractions/{path.name}.md", "")

    def docx(self):
        from docx import Document
        path = self.root / "raw/assets/example.docx"
        doc = Document(); doc.add_heading("Document title unique", 0)
        doc.add_paragraph("Meaningful evidence survives extraction with its source identity. " * 20)
        doc.save(path)
        return path

    def pptx(self, hidden=False, blank=False, broken=False):
        from pptx import Presentation
        from pptx.util import Inches
        path = self.root / "raw/assets/example.pptx"
        deck = Presentation()
        for i in range(1,4):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            if not (blank and i == 2):
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text = f"Slide evidence {i}. " + "Readable content for citation. " * 15
            slide.notes_slide.notes_text_frame.text = f"Speaker note only {i}\vSoft break preserved {i}"
            if hidden and i == 2:slide._element.set("show", "0")
        if broken:
            stream = io.BytesIO();deck.save(stream)
            with zipfile.ZipFile(stream) as src, zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as dst:
                for info in src.infolist():
                    dst.writestr(info, b"broken XML" if info.filename == "ppt/slides/slide2.xml" else src.read(info.filename))
        else:deck.save(path)
        return path

    def require_native(self):
        if not C.backend_available("anydoc", site=str(SITE)):
            self.skipTest("Pinned local AnyDoc wheel not installed; native quality not tested")

    def test_auto_route_never_chooses_anydoc(self):
        with patch.object(C,"backend_available",return_value=True):
            self.assertEqual(C.resolve_backend("auto",".docx"),"pandoc")
            self.assertEqual(C.resolve_backend("auto",".pptx"),"markitdown")
        with patch.object(C,"backend_available",return_value=False):
            self.assertEqual(C.resolve_backend("auto",".pptx"),"builtin")

    def test_unsupported_formats_do_not_fall_through(self):
        for ext in [".pdf",".xlsx",".csv"]:
            p=self.root/('test'+ext);p.write_bytes(b'not opened')
            with self.assertRaisesRegex(RuntimeError,"only .pptx and .docx"):
                C.convert_item(self.item(p),self.args())

    def test_cli_requires_selection(self):
        self.assertEqual(C.main(["--backend","anydoc","--convert"]),2)

    def test_missing_dependency_leaves_no_sidecar(self):
        p=self.docx();args=self.args(anydoc_site=str(self.root/'missing'))
        count,failures,deferred=C.write_conversions([self.item(p)],self.root,args)
        self.assertEqual(count,0);self.assertEqual(len(failures),1);self.assertFalse(deferred)
        self.assertFalse((self.root/self.item(p).target).exists())

    def test_empty_manifest_never_expands_to_whole_root(self):
        manifest=self.root/'empty.txt';manifest.write_text('')
        with patch.object(C.Path,'cwd',return_value=self.root):
            result=C.main(['--backend','anydoc','--convert','--include-manifest',str(manifest),'--roots','raw/assets'])
        self.assertEqual(result,2)

    def test_pilot_blank_slide_is_review_not_ok(self):
        self.require_native();p=self.pptx(blank=True)
        args=self.args(pilot=True,pilot_backends='anydoc',output_dir='wiki/_outputs/pilot')
        self.assertEqual(C.run_pilot([self.item(p)],self.root,args),0)
        report=self.root/'wiki/_outputs/pilot/extraction-pilot/pilot-comparison.md'
        self.assertIn('review(',report.read_text(encoding='utf-8'))

    def test_wrong_version_is_rejected_before_import(self):
        site=self.root/'site';dist=site/'firecrawl_anydoc-9.9.dist-info';dist.mkdir(parents=True)
        (dist/'METADATA').write_text('Name: firecrawl-anydoc\nVersion: 9.9\n')
        with self.assertRaisesRegex(RuntimeError,'requires isolated'):
            C.convert_item(self.item(self.docx()),self.args(anydoc_site=str(site)))

    def test_timeout_fails_without_sidecar(self):
        p=self.docx()
        with patch.object(C.subprocess,'run',side_effect=subprocess.TimeoutExpired('worker',300)):
            count,failures,_=C.write_conversions([self.item(p)],self.root,self.args())
        self.assertEqual(count,0);self.assertIn('timed out',failures[0][1])
        self.assertFalse((self.root/self.item(p).target).exists())

    def test_package_relationship_traversal_and_entities_rejected(self):
        with self.assertRaises(RuntimeError):C.anydoc_part('ppt/slides/slide1.xml','../../../outside')
        with self.assertRaises(RuntimeError):C.anydoc_xml(b'<!DOCTYPE a [<!ENTITY b "c">]><a/>')
        with self.assertRaises(RuntimeError):C.anydoc_xml('<!DOCTYPE a [<!ENTITY b "c">]><a/>'.encode('utf-16'))

    def test_docx_title_create_only_and_provenance(self):
        self.require_native();p=self.docx();before=p.read_bytes();item=self.item(p)
        count,failures,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual((count,failures),(1,[]));target=self.root/item.target
        text=target.read_text(encoding='utf-8');self.assertIn('Document title unique',text)
        self.assertIn('anydoc_version: 0.2.4',text);self.assertIn(hashlib.sha256(before).hexdigest(),text)
        self.assertEqual(C.audit_markdown_file(target,self.root).status,'ok')
        first=target.read_bytes();count,_,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual(count,0);self.assertEqual(target.read_bytes(),first);self.assertEqual(p.read_bytes(),before)

    def test_cli_conversion_audit_and_repeat(self):
        self.require_native();p=self.docx();before=p.read_bytes()
        command=[sys.executable,'-X','utf8',str(ROOT/'tools/source-to-markdown.py'),
            '--roots','raw/assets','--backend','anydoc','--anydoc-site',str(SITE),'--convert','--validate',
            '--sidecar','--output-dir','wiki/_extractions','--report-dir','wiki/_outputs/source-conversions',
            '--include-source','raw/assets/example.docx','--strict']
        run=subprocess.run(command,cwd=self.root,capture_output=True,text=True,encoding='utf-8',timeout=30)
        self.assertEqual(run.returncode,0,run.stdout+run.stderr)
        targets=list((self.root/'wiki/_extractions').rglob('*.md'))
        self.assertEqual(len(targets),1);first=targets[0].read_bytes()
        self.assertEqual(C.audit_markdown_file(targets[0],self.root).status,'ok')
        repeated=subprocess.run(command,cwd=self.root,capture_output=True,text=True,encoding='utf-8',timeout=30)
        self.assertEqual(repeated.returncode,0,repeated.stdout+repeated.stderr)
        self.assertEqual(targets[0].read_bytes(),first);self.assertEqual(p.read_bytes(),before)

    def test_pptx_notes_hidden_order_and_audit(self):
        self.require_native();p=self.pptx(hidden=True);before=p.read_bytes();item=self.item(p)
        count,failures,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual((count,failures),(1,[]));target=self.root/item.target;text=target.read_text(encoding='utf-8')
        for i in range(1,4):
            self.assertIn(f'## Slide {i}',text);self.assertEqual(text.count(f'Speaker note only {i}'),1)
            self.assertIn(f'> Soft break preserved {i}',text)
        self.assertEqual(text.count('### Speaker notes'),3)
        self.assertIn('anydoc_hidden_slides: 2',text);self.assertIn('_Hidden slide',text)
        self.assertEqual(C.audit_markdown_file(target,self.root).status,'ok')
        target.write_text(text.replace('## Slide 2','## Slide 7'),encoding='utf-8')
        self.assertIn('anydoc-slide-coverage-invalid',C.audit_markdown_file(target,self.root).issues)
        self.assertEqual(p.read_bytes(),before)

    def test_blank_slide_requires_review(self):
        self.require_native();p=self.pptx(blank=True);item=self.item(p)
        count,failures,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual((count,failures),(1,[]))
        self.assertIn('anydoc-empty-slide-review-required',C.audit_markdown_file(self.root/item.target,self.root).issues)

    def test_broken_slide_never_publishes_partial_success(self):
        self.require_native();p=self.pptx(broken=True);before=p.read_bytes();item=self.item(p)
        count,failures,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual(count,0);self.assertEqual(len(failures),1)
        self.assertFalse((self.root/item.target).exists());self.assertEqual(p.read_bytes(),before)

    def test_source_change_detected_at_publication(self):
        p=self.docx();item=self.item(p)
        def changing(*args):
            p.write_bytes(b'changed by test');return 'some text','note','anydoc'
        with patch.object(C,'convert_item',side_effect=changing):
            count,failures,_=C.write_conversions([item],self.root,self.args())
        self.assertEqual(count,0);self.assertIn('source changed',failures[0][1]);self.assertFalse((self.root/item.target).exists())


if __name__ == '__main__':unittest.main()
